import re

from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from .config import (
    COR_DATA,
    COR_LINHA,
    COR_TEXTO,
    COR_TITULO,
    FONTE_TEXTO,
    FONTE_TITULO,
    MARGEM_DIREITA,
    MARGEM_ESQUERDA,
    MARGEM_INFERIOR,
    MARGEM_SUPERIOR,
    RECUO_BULLET_MARCADOR,
    RECUO_BULLET_TEXTO,
    TAM_DATA,
    TAM_ROTULO,
    TAM_SECAO,
    TAM_TEXTO,
)
from .conteudo import NOME

PADRAO_LINK = re.compile(
    r"(https?://\S*[\w/-]|(?:github|linkedin)\.com/\S*[\w/-]|[\w.+-]+@[\w-]+(?:\.[\w-]+)+)"
)


def configurar_documento(doc):
    for section in doc.sections:
        section.top_margin = Inches(MARGEM_SUPERIOR)
        section.bottom_margin = Inches(MARGEM_INFERIOR)
        section.left_margin = Inches(MARGEM_ESQUERDA)
        section.right_margin = Inches(MARGEM_DIREITA)

    normal = doc.styles["Normal"]
    normal.font.name = FONTE_TEXTO
    normal.font.size = Pt(TAM_TEXTO)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0

    doc.core_properties.author = NOME
    doc.core_properties.title = f"Currículo {NOME}"


def set_run_font(run, size=TAM_TEXTO, bold=False, italic=False, color=COR_TEXTO, font=FONTE_TEXTO, underline=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.underline = underline
    run.font.name = font
    run.font.color.rgb = RGBColor.from_string(color)
    rFonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:cs"), font)
    rFonts.set(qn("w:eastAsia"), font)


def url_de(trecho):
    if "@" in trecho and not trecho.startswith("http"):
        return "mailto:" + trecho
    if trecho.startswith("http"):
        return trecho
    return "https://" + trecho


def add_hyperlink(paragraph, url, texto, size, font, color, bold=False, italic=False):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = paragraph.add_run(texto)
    set_run_font(run, size=size, bold=bold, italic=italic, color=color, font=font, underline=True)
    paragraph._p.remove(run._r)
    hyperlink.append(run._r)
    paragraph._p.append(hyperlink)


def add_texto(paragraph, texto, size=TAM_TEXTO, font=FONTE_TEXTO, color=COR_TEXTO, bold=False, italic=False, linkar=True):
    partes = PADRAO_LINK.split(texto) if linkar else [texto]
    for i, parte in enumerate(partes):
        if not parte:
            continue
        if linkar and i % 2 == 1:
            add_hyperlink(paragraph, url_de(parte), parte, size, font, color, bold, italic)
        else:
            run = paragraph.add_run(parte)
            set_run_font(run, size=size, bold=bold, italic=italic, color=color, font=font)


def add_rotulo(paragraph, texto):
    run = paragraph.add_run(texto)
    set_run_font(run, size=TAM_ROTULO, bold=True, font=FONTE_TITULO)


def add_borda_inferior(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), COR_LINHA)
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_heading_section(doc, texto):
    p = doc.add_paragraph()
    add_borda_inferior(p)
    run = p.add_run(texto.upper())
    set_run_font(run, size=TAM_SECAO, bold=True, color=COR_TITULO, font=FONTE_TITULO)
    p.paragraph_format.space_before = Pt(7)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    return p


def add_periodo(doc, texto, space_after=0, keep_with_next=False):
    p = doc.add_paragraph()
    run = p.add_run(texto)
    set_run_font(run, size=TAM_DATA, italic=True, color=COR_DATA, font=FONTE_TITULO)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = keep_with_next
    return p


def add_bullet(doc, texto):
    p = doc.add_paragraph(style="List Bullet")
    add_texto(p, texto)
    p.paragraph_format.left_indent = Inches(RECUO_BULLET_TEXTO)
    p.paragraph_format.first_line_indent = Inches(-RECUO_BULLET_MARCADOR)
    p.paragraph_format.space_after = Pt(0)
    return p
