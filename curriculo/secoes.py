import re

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .config import (
    COR_CONTATO,
    COR_TITULO,
    FONTE_TITULO,
    TAM_CONTATO,
    TAM_NOME,
)
from .conteudo import CONTATO, EXPERIENCIAS, FORMACAO, IDIOMAS, NOME, RESUMO, SUPORTE_TI
from .dados import filtrar_cursos, filtrar_projetos, formatar_techs
from .formatacao import (
    add_bullet,
    add_heading_section,
    add_periodo,
    add_rotulo,
    add_texto,
    set_run_font,
)


def add_cabecalho(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(NOME)
    set_run_font(run, size=TAM_NOME, bold=True, color=COR_TITULO, font=FONTE_TITULO)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for i, parte in enumerate(CONTATO.split(" | ")):
        if i:
            run = p.add_run("  |  ")
            set_run_font(run, size=TAM_CONTATO, color=COR_CONTATO)
        add_texto(p, parte, size=TAM_CONTATO, color=COR_CONTATO, linkar="@" in parte)
    p.paragraph_format.space_after = Pt(2)


def add_objetivo(doc, cargo, categoria, nivel, objetivo):
    add_heading_section(doc, "Objetivo Profissional")
    p = doc.add_paragraph()
    add_rotulo(p, f"{cargo} {categoria} {nivel}")
    p.paragraph_format.space_after = Pt(1)
    p = doc.add_paragraph()
    add_texto(p, objetivo)


def add_conhecimentos(doc, techs, suporte):
    add_heading_section(doc, "Conhecimentos Técnicos")
    if suporte:
        p = doc.add_paragraph()
        add_rotulo(p, "Suporte e TI: ")
        add_texto(p, SUPORTE_TI)
    if techs:
        p = doc.add_paragraph()
        add_rotulo(p, "Desenvolvimento: ")
        add_texto(p, formatar_techs(techs))
    p = doc.add_paragraph()
    add_rotulo(p, "Idiomas: ")
    add_texto(p, " • ".join(IDIOMAS))


def add_resumo(doc):
    add_heading_section(doc, "Resumo de Qualificações")
    p = doc.add_paragraph()
    add_texto(p, RESUMO)


def add_formacao(doc):
    add_heading_section(doc, "Formação Acadêmica")
    for i, item in enumerate(FORMACAO):
        p = doc.add_paragraph()
        add_rotulo(p, item["titulo"])
        p.paragraph_format.space_before = Pt(0 if i == 0 else 3)
        p.paragraph_format.keep_with_next = True
        add_periodo(doc, item["periodo"])


def add_experiencia(doc):
    add_heading_section(doc, "Experiência Profissional")
    for i, exp in enumerate(EXPERIENCIAS):
        p = doc.add_paragraph()
        add_rotulo(p, exp["titulo"])
        p.paragraph_format.space_before = Pt(0 if i == 0 else 5)
        p.paragraph_format.keep_with_next = True
        add_periodo(doc, exp["periodo"], space_after=1, keep_with_next=True)
        for bullet in exp["bullets"]:
            add_bullet(doc, bullet)


def add_cursos(doc, categorias_lista, suporte):
    cursos = filtrar_cursos(categorias_lista, suporte)
    if not cursos:
        return
    add_heading_section(doc, "Cursos / Certificados")
    for c in cursos:
        p = doc.add_paragraph()
        partes = c["titulo"].split(" | ", 1)
        add_rotulo(p, partes[0])
        resto = " | " + (partes[1] + " | " if len(partes) > 1 else "") + c["data"]
        add_texto(p, resto)
        p.paragraph_format.space_after = Pt(1)


def add_projetos(doc, categorias_lista):
    projetos = filtrar_projetos(categorias_lista)
    if not projetos:
        return
    add_heading_section(doc, "Projetos")
    for pjt in projetos:
        p = doc.add_paragraph()
        m = re.match(r"^(.*?)(\s*\(.*\))?$", pjt["nome"])
        nome = m.group(1)
        complemento = m.group(2) or ""
        add_texto(p, nome, bold=True, linkar=False)
        if complemento:
            add_texto(p, complemento, linkar=False)
        add_texto(p, f" | {pjt['descricao']} Publicado no GitHub: {pjt['link']}")
        p.paragraph_format.space_after = Pt(2)
